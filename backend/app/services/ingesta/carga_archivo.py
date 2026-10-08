"""Cargar el Excel del Z-tracking subido desde la interfaz — `US-58`.

    «Como usuario de Compras, quiero subir el Excel del Z-tracking desde TrackIn
    y ver qué entró y qué no, para mantener los pedidos al día sin depender de
    un desarrollador.»

**La carga es la misma de siempre.** Este módulo no lee ni normaliza nada: deja
el archivo en disco, lo lee con `FuenteZTracking` y lo pasa por `carga.cargar`
(`US-31`, `US-32`, con su recálculo de `US-12`). Lo nuevo son tres guardas que
el script no necesitaba, porque quien lo corría sabía qué archivo le daba:

1. **Un archivo equivocado se rechaza entero, antes de tocar ningún pedido.** Lo
   peligroso no es un error de lectura: es un Excel que se lee bien pero no es
   el Z-tracking. Sin las hojas `PRODUCCION` e `IDA` el lector no encontraría
   ninguna línea y la carga, de buena fe, marcaría **todos** los pedidos como
   ausentes. Por eso se exigen las dos hojas y al menos una línea legible.
2. **Una carga a la vez.** Compras usa una cuenta compartida (30/09): dos
   personas pueden subir el archivo al mismo tiempo. Un candado de PostgreSQL
   hace que la segunda reciba «hay una carga en curso» en vez de mezclarse.
3. **Queda registro**, también de lo rechazado (`cargas_ztracking`).
"""

from __future__ import annotations

import datetime as dt
import logging
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path, PureWindowsPath
from typing import Any, Final

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.carga_ztracking import APLICADA, RECHAZADA, CargaZTracking
from app.services.ingesta.carga import ResolutorDestino, ResultadoCarga, cargar
from app.services.ingesta.dto import PedidoCrudo
from app.services.ingesta.informe import InformeValidacion, construir_informe
from app.services.ingesta.ztracking import HOJAS_EN_ALCANCE, FuenteZTracking, HojaInesperada

logger = logging.getLogger(__name__)

#: El WK38 real pesa ~1 MB; 20 MB deja margen sin aceptar cualquier cosa.
TAMANO_MAXIMO_BYTES: Final = 20 * 1024 * 1024
EXTENSIONES: Final = (".xlsx", ".xlsm")

#: Clave del candado de PostgreSQL que serializa las cargas. Arbitraria, fija.
_CANDADO_CARGA: Final = 5_801_058


class ArchivoInvalido(ValueError):
    """El archivo no es un Z-tracking utilizable. No se tocó ningún pedido."""


class CargaEnCurso(RuntimeError):
    """Otra carga está corriendo; esta no empezó."""


@dataclass(frozen=True, slots=True)
class _FuenteLeida:
    """Las líneas ya leídas, con la forma de `FuentePedidos`.

    El archivo se lee una vez, antes de abrir la carga, para poder rechazarlo
    entero si no trae nada. Esto evita leerlo dos veces.
    """

    nombre: str
    descripcion: str
    pedidos: list[PedidoCrudo]

    async def obtener_pedidos(self) -> list[PedidoCrudo]:
        return self.pedidos


def _sin_ruta(nombre_archivo: str | None) -> str:
    """El nombre del archivo sin la carpeta del equipo de quien lo subió.

    `PureWindowsPath` y no `PurePath`: este último sigue las reglas del sistema
    donde corre, y en el servidor Linux no reconocería `\\` como separador, así
    que «C:\\Users\\x\\WK40.xlsx» quedaría entero. `PureWindowsPath` entiende
    `\\` y `/` en cualquier sistema.
    """
    return PureWindowsPath(nombre_archivo or "").name.strip()


def validar_nombre(nombre_archivo: str) -> str:
    """El nombre sin ruta, y solo si es un Excel."""
    nombre = _sin_ruta(nombre_archivo)
    if not nombre:
        raise ArchivoInvalido("El archivo no tiene nombre.")
    if not nombre.lower().endswith(EXTENSIONES):
        raise ArchivoInvalido(
            f"«{nombre}» no es un Excel (.xlsx). Exporte el Z-tracking de SAP como libro de Excel."
        )
    return nombre[:255]


def _verificar_hojas(ruta: Path, nombre: str) -> None:
    try:
        libro = load_workbook(ruta, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, OSError, KeyError) as exc:
        raise ArchivoInvalido(
            f"«{nombre}» no se pudo abrir como Excel: está dañado o no es un .xlsx."
        ) from exc
    try:
        faltan = [hoja for hoja in HOJAS_EN_ALCANCE if hoja not in libro.sheetnames]
    finally:
        libro.close()
    if faltan:
        raise ArchivoInvalido(
            f"«{nombre}» no parece el Z-tracking: le faltan las hojas {', '.join(faltan)}. "
            "No se cargó nada, para no dar por ausentes los pedidos que no vienen."
        )


async def _leer(contenido: bytes, nombre: str) -> tuple[list[PedidoCrudo], FuenteZTracking]:
    # `delete=False` y borrado a mano: en Windows un temporal abierto no se
    # puede volver a abrir por nombre, que es lo que hace openpyxl.
    with tempfile.NamedTemporaryFile(suffix=Path(nombre).suffix, delete=False) as tmp:
        tmp.write(contenido)
        ruta = Path(tmp.name)
    try:
        _verificar_hojas(ruta, nombre)
        fuente = FuenteZTracking(ruta)
        try:
            pedidos = await fuente.obtener_pedidos()
        except HojaInesperada as exc:
            raise ArchivoInvalido(str(exc).replace(ruta.name, nombre)) from exc
    finally:
        ruta.unlink(missing_ok=True)

    if not pedidos:
        raise ArchivoInvalido(
            f"«{nombre}» tiene las hojas del Z-tracking pero ninguna línea legible. "
            "No se cargó nada, para no dar por ausentes todos los pedidos."
        )
    return pedidos, fuente


def _avisos(resultado: ResultadoCarga) -> list[dict[str, str]]:
    """Lo que entró bien pero conviene corregir en SAP (`US-52`, `US-54`)."""
    return [
        {
            "clave": f"{linea.oc_numero}-{linea.posicion_oc}",
            "tipo": linea.motivo,
            "detalle": linea.detalle,
        }
        for linea in (*resultado.destinos_de_la_fuente, *resultado.destinos_discrepantes)
    ]


def _registro(
    *,
    id_usuario: int,
    archivo: str,
    tamano: int,
    instante: dt.datetime,
    informe: InformeValidacion | None = None,
    avisos: list[dict[str, str]] | None = None,
    motivo_rechazo: str | None = None,
) -> CargaZTracking:
    fila = CargaZTracking(
        id_usuario=id_usuario,
        archivo=archivo,
        tamano_bytes=tamano,
        realizada_en=instante,
        estado=RECHAZADA if motivo_rechazo else APLICADA,
        motivo_rechazo=motivo_rechazo[:500] if motivo_rechazo else None,
    )
    if informe is not None:
        datos: dict[str, Any] = informe.como_dict()
        datos["avisos"] = avisos or []
        fila.informe = datos
        fila.recibidas = informe.recibidas
        fila.insertadas = informe.insertadas
        fila.actualizadas = informe.actualizadas
        fila.sin_cambios = informe.sin_cambios
        fila.ausentes = informe.ausentes
        fila.no_entraron = informe.no_entraron
        fila.entraron_sin_rastreo = informe.entraron_sin_rastreo
    return fila


async def cargar_archivo(
    sesion: AsyncSession,
    *,
    contenido: bytes,
    nombre_archivo: str,
    id_usuario: int,
    resolutor_destino: ResolutorDestino | None = None,
    instante: dt.datetime | None = None,
) -> CargaZTracking:
    """Carga el Excel subido y deja su registro. **No hace commit.**

    Lanza `ArchivoInvalido` sin haber tocado la base (quien llama registra el
    rechazo con `registrar_rechazo`, en su propia transacción) y `CargaEnCurso`
    si otra carga tiene el candado.
    """
    nombre = validar_nombre(nombre_archivo)
    if not contenido:
        raise ArchivoInvalido(f"«{nombre}» está vacío.")
    if len(contenido) > TAMANO_MAXIMO_BYTES:
        raise ArchivoInvalido(
            f"«{nombre}» pesa {len(contenido) // (1024 * 1024)} MB; el máximo es "
            f"{TAMANO_MAXIMO_BYTES // (1024 * 1024)} MB."
        )

    pedidos, fuente = await _leer(contenido, nombre)

    # Candado de transacción: se suelta solo con el commit o el rollback.
    libre = await sesion.scalar(select(text(f"pg_try_advisory_xact_lock({_CANDADO_CARGA})")))
    if not libre:
        raise CargaEnCurso(
            "Hay otra carga del Z-tracking en curso. Espere a que termine y vuelva a intentar."
        )

    ahora = instante or dt.datetime.now(dt.UTC)
    resultado = await cargar(
        sesion,
        _FuenteLeida("ztracking", f"Z-tracking subido ({nombre})", pedidos),
        resolutor_destino=resolutor_destino,
    )
    informe = construir_informe("ztracking", resultado, fuente.ilegibles)
    fila = _registro(
        id_usuario=id_usuario,
        archivo=nombre,
        tamano=len(contenido),
        instante=ahora,
        informe=informe,
        avisos=_avisos(resultado),
    )
    sesion.add(fila)
    await sesion.flush()
    logger.info(
        "Carga del Z-tracking %r por el usuario %s: %s",
        nombre,
        id_usuario,
        informe.como_texto(ejemplos=0).splitlines()[1].strip(),
    )
    return fila


def registrar_rechazo(
    sesion: AsyncSession,
    *,
    nombre_archivo: str,
    tamano: int,
    id_usuario: int,
    motivo: str,
    instante: dt.datetime | None = None,
) -> CargaZTracking:
    """Deja constancia de un archivo rechazado. **No hace commit.**"""
    nombre = _sin_ruta(nombre_archivo)[:255] or "(sin nombre)"
    fila = _registro(
        id_usuario=id_usuario,
        archivo=nombre,
        tamano=tamano,
        instante=instante or dt.datetime.now(dt.UTC),
        motivo_rechazo=motivo,
    )
    sesion.add(fila)
    return fila


__all__ = [
    "EXTENSIONES",
    "TAMANO_MAXIMO_BYTES",
    "ArchivoInvalido",
    "CargaEnCurso",
    "cargar_archivo",
    "registrar_rechazo",
    "validar_nombre",
]
